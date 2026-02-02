"""
MCP Server: Ingestion Agent
Handles file upload and initial processing using FastMCP.
"""

from fastmcp import FastMCP
from pathlib import Path
import PyPDF2
from PIL import Image
import openpyxl
from docx import Document as DocxDocument
import json

mcp = FastMCP("Ingestion Agent")

@mcp.tool()
def process_pdf(file_path: str) -> str:
    """
    Extract text and metadata from PDF file.
    Returns JSON with pages, text, and metadata.
    """
    try:
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            
            pages_data = []
            for page_num, page in enumerate(reader.pages, start=1):
                text = page.extract_text()
                text_density = len(text.strip()) / max(len(text), 1)
                
                pages_data.append({
                    "page_number": page_num,
                    "text": text,
                    "text_density": text_density,
                    "needs_ocr": text_density < 0.3
                })
            
            result = {
                "file_type": "pdf",
                "total_pages": len(reader.pages),
                "pages": pages_data,
                "metadata": {
                    "file_name": Path(file_path).name,
                    "total_pages": len(reader.pages)
                }
            }
            
            return json.dumps(result)
            
    except Exception as e:
        return json.dumps({"error": str(e)})


@mcp.tool()
def process_image(file_path: str) -> str:
    """
    Process image file and return metadata.
    Returns JSON indicating OCR is needed.
    """
    try:
        img = Image.open(file_path)
        
        result = {
            "file_type": "image",
            "total_pages": 1,
            "pages": [{
                "page_number": 1,
                "text": "",
                "text_density": 0.0,
                "needs_ocr": True
            }],
            "metadata": {
                "file_name": Path(file_path).name,
                "image_size": img.size,
                "format": img.format
            }
        }
        
        return json.dumps(result)
        
    except Exception as e:
        return json.dumps({"error": str(e)})


@mcp.tool()
def process_excel(file_path: str) -> str:
    """
    Extract data from Excel file.
    Returns JSON with sheet data.
    """
    try:
        wb = openpyxl.load_workbook(file_path)
        
        sheets_data = []
        for sheet_idx, sheet in enumerate(wb.worksheets, start=1):
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append([str(cell) if cell is not None else "" for cell in row])
            
            content = "\n".join(["\t".join(row) for row in rows])
            
            sheets_data.append({
                "page_number": sheet_idx,
                "text": content,
                "text_density": 1.0,
                "needs_ocr": False,
                "sheet_name": sheet.title
            })
        
        result = {
            "file_type": "excel",
            "total_pages": len(sheets_data),
            "pages": sheets_data,
            "metadata": {
                "file_name": Path(file_path).name,
                "total_sheets": len(sheets_data)
            }
        }
        
        return json.dumps(result)
        
    except Exception as e:
        return json.dumps({"error": str(e)})


@mcp.tool()
def process_word(file_path: str) -> str:
    """
    Extract text from Word document.
    Returns JSON with document text.
    """
    try:
        doc = DocxDocument(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        content = "\n".join(paragraphs)
        
        result = {
            "file_type": "word",
            "total_pages": 1,
            "pages": [{
                "page_number": 1,
                "text": content,
                "text_density": 1.0,
                "needs_ocr": False
            }],
            "metadata": {
                "file_name": Path(file_path).name,
                "paragraph_count": len(paragraphs)
            }
        }
        
        return json.dumps(result)
        
    except Exception as e:
        return json.dumps({"error": str(e)})


if __name__ == "__main__":
    import os
    from dotenv import load_dotenv
    load_dotenv()
    
    MCP_INGESTION_PORT = int(os.getenv("MCP_INGESTION_PORT", 8001))
    
    mcp.run(transport="http", host="127.0.0.1", port=MCP_INGESTION_PORT)
